
import base64
import io
import unicodedata
import re
from difflib import SequenceMatcher
import openpyxl
from odoo import api, fields, models, _
from odoo.exceptions import UserError

def normalize_name(name):
    if not name:
        return ""
    name = str(name).replace('\xa0', ' ').replace('\u200b', ' ').replace('\t', ' ')
    name = ''.join(c for c in unicodedata.normalize('NFD', name) if unicodedata.category(c) != 'Mn')
    name = name.upper()
    name = re.sub(r'[^A-Z0-9\s]', ' ', name)
    name = re.sub(r'\s+', ' ', name).strip()
    return name

def token_sort_ratio(a, b):
    a_tokens = sorted(normalize_name(a).split())
    b_tokens = sorted(normalize_name(b).split())
    return SequenceMatcher(None, ' '.join(a_tokens), ' '.join(b_tokens)).ratio()

class SgsPerdiemDepositImportWizard(models.TransientModel):
    _name = 'sgs.perdiem.deposit.import.wizard'
    _description = 'Importador inteligente de depositos'

    file = fields.Binary('Archivo Excel', required=True)
    filename = fields.Char('Nombre archivo')
    date_default = fields.Date('Fecha por defecto', default=fields.Date.context_today)
    line_ids = fields.One2many('sgs.perdiem.deposit.import.line', 'wizard_id', string='Lineas a conciliar')
    state = fields.Selection([('draft','Carga'),('to_conciliate','Conciliar'),('done','Hecho')], default='draft')

    def action_parse_file(self):
        self.ensure_one()
        data = base64.b64decode(self.file)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        # Buscar header real (primera fila con CUSTODIO)
        header_row_idx = 1
        col_map = {}
        for r in range(1, 10):
            row_vals = [str(c.value or '') for c in ws[r]]
            norm_row = [normalize_name(v) for v in row_vals]
            # Detecta encabezado
            if any('CUSTODIO' in n for n in norm_row):
                header_row_idx = r
                for idx, h in enumerate(norm_row):
                    if 'FECHA' in h and 'DEPOSITO' in h:
                        col_map['date'] = idx
                    elif 'CUSTODIO' in h:
                        col_map['custodio'] = idx
                    elif h == 'MONTO' or 'MONTO' in h:
                        col_map['amount'] = idx
                    elif 'CONCEPTO' in h:
                        col_map['concept'] = idx
                break
        
        if 'custodio' not in col_map:
            raise UserError(f"No encontre columna CUSTODIO. Encabezados detectados fila {header_row_idx}: {row_vals}")

        custodians = self.env['sgs.custodian'].search([])
        employees = self.env['hr.employee'].search([])
        cust_map = {normalize_name(c.name): c for c in custodians}
        emp_map = {normalize_name(e.name): e for e in employees}

        lines = []
        for row in ws.iter_rows(min_row=header_row_idx+1, values_only=True):
            if not row or not any(row):
                continue
            row = list(row) + [None]*20
            raw_cust = str(row[col_map.get('custodio',1)] or '').strip()
            if not raw_cust:
                continue
            # Filtro: si el nombre es numerico o dice MONTO TRASPASO, es basura del excel
            if raw_cust.replace('.','',1).isdigit() or 'MONTO' in normalize_name(raw_cust) or len(raw_cust) < 5:
                continue

            raw_amount = row[col_map.get('amount',5)] if 'amount' in col_map else row[5]
            try:
                amount = float(str(raw_amount).replace(',','').replace('$','') or 0)
            except:
                amount = 0

            raw_date = row[col_map.get('date',0)] if 'date' in col_map else row[0]
            date_val = self.date_default
            if raw_date:
                if hasattr(raw_date, 'year'):
                    date_val = raw_date.date() if hasattr(raw_date, 'date') else raw_date
                else:
                    try:
                        from dateutil import parser as date_parser
                        date_val = date_parser.parse(str(raw_date), dayfirst=True).date()
                    except:
                        date_val = self.date_default

            concept = str(row[col_map.get('concept',4)] if 'concept' in col_map else (row[4] or 'Deposito semanal viaticos')).strip()

            norm = normalize_name(raw_cust)
            custodian = cust_map.get(norm)
            employee = emp_map.get(norm)
            status = 'matched'
            score = 1.0

            if not custodian and employee:
                custodian = self.env['sgs.custodian'].search([('employee_id','=',employee.id)], limit=1)
            
            if not custodian:
                best_match = None
                best_score = 0
                for c_name, c_rec in cust_map.items():
                    # debe compartir al menos 1 token para no comparar todo
                    if not set(norm.split()) & set(c_name.split()):
                        continue
                    s = token_sort_ratio(norm, c_name)
                    if s > best_score:
                        best_score = s
                        best_match = c_rec
                if best_score >= 0.85:
                    custodian = best_match
                    score = best_score
                    status = 'matched_auto'
                elif best_score >= 0.55:
                    custodian = best_match
                    score = best_score
                    status = 'to_conciliate'
                else:
                    status = 'not_found'
                    score = best_score

            lines.append((0,0,{
                'custodio_raw': raw_cust,
                'custodio_normalized': norm,
                'custodian_id': custodian.id if custodian else False,
                'employee_id': employee.id if employee else (custodian.employee_id.id if custodian and custodian.employee_id else False),
                'date': date_val,
                'amount': amount,
                'concept': concept,
                'match_score': score,
                'state': status,
                'to_import': bool(custodian and amount > 0),
            }))
        
        if not lines:
            raise UserError(f"No se detectaron lineas validas. Mapa columnas: {col_map}. Revisa que la columna CUSTODIO tenga nombres.")

        self.line_ids = [(5,0,0)] + lines
        self.state = 'to_conciliate'
        return {
            'type':'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode':'form',
            'target':'new',
        }

    def action_create_deposits(self):
        to_create = self.line_ids.filtered(lambda l: l.custodian_id and l.to_import and l.amount > 0)
        if not to_create:
            raise UserError("No hay lineas validas para importar. Asigna el custodio manualmente en las lineas rojas.")
        vals_list = []
        for line in to_create:
            vals_list.append({
                'custodian_id': line.custodian_id.id,
                'date': line.date,
                'amount': line.amount,
                'concept': line.concept,
            })
        deposits = self.env['sgs.perdiem.deposit'].create(vals_list)
        self.state = 'done'
        return {
            'type':'ir.actions.act_window',
            'name':'Depositos creados',
            'res_model':'sgs.perdiem.deposit',
            'domain':[('id','in',deposits.ids)],
            'view_mode':'list,form',
        }

class SgsPerdiemDepositImportLine(models.TransientModel):
    _name = 'sgs.perdiem.deposit.import.line'
    _description = 'Linea de conciliacion de deposito'
    wizard_id = fields.Many2one('sgs.perdiem.deposit.import.wizard', required=True, ondelete='cascade')
    custodio_raw = fields.Char('Nombre en Excel', readonly=True)
    custodio_normalized = fields.Char('Normalizado', readonly=True)
    employee_id = fields.Many2one('hr.employee', string='Empleado encontrado')
    custodian_id = fields.Many2one('sgs.custodian', string='Custodio a depositar')
    date = fields.Date('Fecha')
    amount = fields.Float('Monto')
    concept = fields.Char('Concepto')
    match_score = fields.Float('Score', digits=(3,2))
    state = fields.Selection([
        ('matched','Coincidencia exacta'),
        ('matched_auto','Auto-conciliado'),
        ('to_conciliate','Requiere revision'),
        ('not_found','No encontrado')
    ], default='to_conciliate')
    to_import = fields.Boolean('Importar', default=True)
